"""Generate professional, publication-quality DOCX documents from Markdown files
using python-docx with custom typography, corporate palettes, and styled tables.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# Palette
COLOR_PRIMARY = RGBColor(0x1B, 0x36, 0x5D)    # Deep Navy
COLOR_SECONDARY = RGBColor(0x2E, 0x6B, 0x6B)  # Slate Teal
COLOR_TEXT = RGBColor(0x24, 0x29, 0x2F)       # Charcoal text
COLOR_MUTED = RGBColor(0x57, 0x60, 0x6A)      # Muted gray
COLOR_CODE = RGBColor(0x05, 0x50, 0xAE)       # Blue code


def set_cell_background(cell, fill_hex: str):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, value in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{name}')
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_formatted_run(p, text: str, default_font="Segoe UI", default_size=10.5, default_color=COLOR_TEXT):
    # Parses bold **text** and code `code`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(default_size - 0.5)
            run.font.color.rgb = COLOR_CODE
        else:
            run = p.add_run(token)
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color


def markdown_to_docx(md_path: Path, docx_path: Path, title_override: str | None = None):
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        
        # Determine cols
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        for row_idx, r_data in enumerate(table_rows):
            is_header = (row_idx == 0)
            row = tbl.rows[row_idx]
            for col_idx in range(num_cols):
                cell = row.cells[col_idx]
                text = r_data[col_idx] if col_idx < len(r_data) else ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                
                if is_header:
                    set_cell_background(cell, "1B365D")
                    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                    add_formatted_run(p, text, default_font="Segoe UI Semibold", default_size=10, default_color=RGBColor(0xFF, 0xFF, 0xFF))
                else:
                    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    set_cell_margins(cell, top=110, bottom=110, left=160, right=160)
                    add_formatted_run(p, text, default_font="Segoe UI", default_size=9.5, default_color=COLOR_TEXT)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    def flush_code():
        nonlocal code_lines, in_code_block
        if not code_lines:
            in_code_block = False
            return
        
        # Add code box table
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        code_text = "\n".join(code_lines)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()

        # Handle code fences
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle markdown tables
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped[1:-1].split("|")]
            # Check if separator row (---)
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        else:
            if in_table:
                flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            title_text = title_override or stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title_text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:].strip())
            run.font.name = "Segoe UI Semibold"
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:].strip())
            run.font.name = "Segoe UI Semibold"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            continue

        if stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = COLOR_MUTED
            continue

        # Callout blockquote
        if stripped.startswith("> "):
            callout_text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
            
            # Left border styling
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
            tcPr.append(borders)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, callout_text, default_size=10, default_color=RGBColor(0x1B, 0x36, 0x5D))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # Bullet list item
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, stripped[2:].strip())
            continue

        # Numbered list item
        m_num = re.match(r"^(\d+[\.\)])\s+(.*)", stripped)
        if m_num:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_run(p, m_num.group(2).strip())
            continue

        # Standard paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        add_formatted_run(p, stripped)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path} ({docx_path.stat().st_size:,} bytes)")


def main():
    if len(sys.argv) > 1:
        md_file = Path(sys.argv[1])
        if len(sys.argv) > 2:
            docx_file = Path(sys.argv[2])
        else:
            docx_file = md_file.with_suffix(".docx")
        if not md_file.exists():
            print(f"Error: {md_file} does not exist.")
            sys.exit(1)
        markdown_to_docx(md_file, docx_file)
        return

    docs_dir = Path("docs")
    docs_to_convert = [
        ("05_CURRENT_FEATURES_DETAILED_ANALYSIS.md", "05_CURRENT_FEATURES_DETAILED_ANALYSIS.docx"),
        ("06_FLAWS_LIMITATIONS_AND_VULNERABILITIES.md", "06_FLAWS_LIMITATIONS_AND_VULNERABILITIES.docx"),
        ("07_ENTERPRISE_GRADE_TRANSFORMATION_ROADMAP.md", "07_ENTERPRISE_GRADE_TRANSFORMATION_ROADMAP.docx"),
    ]

    for md_name, docx_name in docs_to_convert:
        md_file = docs_dir / md_name
        docx_file = docs_dir / docx_name
        if not md_file.exists():
            print(f"Error: {md_file} does not exist.")
            continue
        markdown_to_docx(md_file, docx_file)


if __name__ == "__main__":
    main()

